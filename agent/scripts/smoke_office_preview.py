"""Real runtime-image smoke: Office rendering, native edit, and refreshed PDF."""
from __future__ import annotations

import asyncio
import base64
import io
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from pptx import Presentation
from pypdf import PdfReader

from helpudoc_agent.office_preview import OfficePreviewService


async def main() -> None:
    service = OfficePreviewService()
    document = Document()
    document.add_paragraph("Office preview smoke")
    document.add_page_break()
    document.add_paragraph("Original second page")
    stream = io.BytesIO()
    document.save(stream)
    source = stream.getvalue()
    first = await service.preview("smoke.docx", source)
    assert len(PdfReader(io.BytesIO(base64.b64decode(first["pdf"]))).pages) == 2
    updated = await service.edit("smoke.docx", source, first["revision"], {
        "paragraphId": "p:0", "start": 0, "end": 6, "quote": "Office",
        "action": "replaceText", "value": "Revised",
    })
    refreshed = await service.preview("smoke.docx", base64.b64decode(updated["content"]))
    pages = PdfReader(io.BytesIO(base64.b64decode(refreshed["pdf"]))).pages
    assert len(pages) == 2 and "Revised preview smoke" in pages[0].extract_text()
    assert "Original second page" in pages[1].extract_text()

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "PowerPoint preview smoke"
    stream = io.BytesIO()
    presentation.save(stream)
    deck = await service.preview("smoke.pptx", stream.getvalue())
    slides = PdfReader(io.BytesIO(base64.b64decode(deck["pdf"]))).pages
    assert len(slides) == 1 and "PowerPoint preview smoke" in slides[0].extract_text()
    print(json.dumps({"office_preview": "ok", "docx_pages": 2, "pptx_slides": 1, "native_edit": "ok"}))


if __name__ == "__main__":
    asyncio.run(main())
