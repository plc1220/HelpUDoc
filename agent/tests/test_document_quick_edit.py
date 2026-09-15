from __future__ import annotations

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from lxml import etree
import pytest

from helpudoc_agent.document_quick_edit import QuickEditError, apply_docx_edit, inspect_docx


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}
STYLES = f'''<w:styles xmlns:w="{W}">
 <w:docDefaults><w:rPrDefault><w:rPr><w:sz w:val="22"/></w:rPr></w:rPrDefault></w:docDefaults>
 <w:style w:type="paragraph" w:styleId="BodyOriginal" w:default="1"><w:name w:val="Original body"/></w:style>
 <w:style w:type="paragraph" w:styleId="UploadedTitle"><w:name w:val="My title"/><w:basedOn w:val="BodyOriginal"/><w:rPr><w:b/><w:sz w:val="40"/></w:rPr></w:style>
 <w:style w:type="character" w:styleId="Emphasis"><w:name w:val="Emphasis"/><w:rPr><w:i/></w:rPr></w:style>
</w:styles>'''.encode()


def package(body: str, *, settings: str = "", extra: dict[str, bytes] | None = None, styles: bytes = STYLES) -> bytes:
    result = BytesIO()
    with ZipFile(result, "w", ZIP_DEFLATED) as archive:
        archive.comment = b"Preserve uploaded metadata"
        archive.writestr("[Content_Types].xml", b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        archive.writestr("word/document.xml", f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{W}" xmlns:custom="urn:custom"><w:body>{body}<w:sectPr><w:pgSz w:w="11906" w:h="16838"/></w:sectPr></w:body></w:document>'''.encode())
        archive.writestr("word/styles.xml", styles)
        archive.writestr("word/settings.xml", f'<w:settings xmlns:w="{W}">{settings}</w:settings>'.encode())
        archive.writestr("word/header1.xml", b'<original header="preserve me"/>')
        archive.writestr("word/media/image1.png", b"unchanged-binary-image")
        for key, value in (extra or {}).items():
            archive.writestr(key, value)
    return result.getvalue()


def edit(**kwargs):
    return {"paragraphId": "p:0", "start": 0, "end": 5, "quote": "Hello", "action": "bold", "value": True, **kwargs}


def content(data: bytes) -> dict[str, bytes]:
    with ZipFile(BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def document(data: bytes):
    return etree.fromstring(content(data)["word/document.xml"])


def test_inspection_inherits_uploaded_styles_and_counts_unicode_and_table_paragraphs():
    data = package('''<w:p><w:pPr><w:pStyle w:val="UploadedTitle"/></w:pPr>
        <w:r><w:t>Hi😀</w:t></w:r><w:r><w:rPr><w:rStyle w:val="Emphasis"/><w:b w:val="0"/></w:rPr><w:t>文</w:t></w:r></w:p>
        <w:tbl><w:tr><w:tc><w:p><w:r><w:t>Cell</w:t></w:r></w:p></w:tc></w:tr></w:tbl>''')
    inspected = inspect_docx(data)
    first, cell = inspected["paragraphs"]
    assert first == {"id": "p:0", "text": "Hi😀文", "styleId": "UploadedTitle", "editable": True, "runs": [
        {"start": 0, "end": 3, "bold": True, "italic": None, "fontSize": 20.0},
        {"start": 3, "end": 4, "bold": False, "italic": True, "fontSize": 20.0},
    ]}
    assert cell["id"] == "p:1" and cell["editable"] and cell["styleId"] == "BodyOriginal"
    assert inspected["styles"] == [{"id": "BodyOriginal", "name": "Original body"}, {"id": "UploadedTitle", "name": "My title"}]


def test_partial_formatting_preserves_upload_parts_and_unrelated_document_xml_bytes():
    untouched = '<w:p custom:token="keep"><w:pPr><w:spacing w:after="87"/></w:pPr><w:r><w:t>Unchanged</w:t></w:r></w:p>'
    data = package('''<w:p custom:id="source"><w:pPr><w:keepNext/></w:pPr><w:r w:rsidR="12"><w:rPr><w:rFonts w:ascii="Brand Font"/><w:color w:val="AABBCC"/><w:shd w:fill="EEDDCC"/></w:rPr><w:t>Hello world</w:t></w:r></w:p>''' + untouched)
    updated = apply_docx_edit(data, edit(start=2, end=8, quote="llo wo"))
    before, after = content(data), content(updated)
    assert {k: v for k, v in before.items() if k != "word/document.xml"} == {k: v for k, v in after.items() if k != "word/document.xml"}
    assert untouched.encode() in after["word/document.xml"]
    assert after["word/document.xml"].startswith(before["word/document.xml"].split(b'<w:p custom:id="source">')[0])
    assert after["word/document.xml"].endswith(before["word/document.xml"].split(untouched.encode())[1])
    with ZipFile(BytesIO(data)) as original, ZipFile(BytesIO(updated)) as result:
        assert result.comment == original.comment
        assert [(x.filename, x.date_time, x.external_attr, x.compress_type) for x in result.infolist()] == [(x.filename, x.date_time, x.external_attr, x.compress_type) for x in original.infolist()]
    runs = document(updated).xpath("//w:body/w:p[1]/w:r", namespaces=NS)
    assert ["".join(r.itertext()) for r in runs] == ["He", "llo wo", "rld"]
    for run in runs:
        assert run.get(f"{{{W}}}rsidR") == "12"
        assert run.find("w:rPr/w:rFonts", NS).get(f"{{{W}}}ascii") == "Brand Font"
        assert run.find("w:rPr/w:shd", NS).get(f"{{{W}}}fill") == "EEDDCC"
    assert runs[0].find("w:rPr/w:b", NS) is None
    assert runs[1].find("w:rPr/w:b", NS).get(f"{{{W}}}val") == "1"
    assert runs[1].find("w:rPr/w:bCs", NS).get(f"{{{W}}}val") == "1"
    assert runs[2].find("w:rPr/w:b", NS) is None


def test_cross_run_replace_preserves_surrounding_style_and_inherits_first_selected_run():
    data = package('''<w:p><w:r><w:rPr><w:i/></w:rPr><w:t>Hello </w:t></w:r><w:r><w:rPr><w:b/></w:rPr><w:t>😀world!</w:t></w:r></w:p>''')
    updated = apply_docx_edit(data, edit(start=3, end=10, quote="lo 😀wor", action="replaceText", value="new text"))
    inspected = inspect_docx(updated)["paragraphs"][0]
    assert inspected["text"] == "Helnew textld!"
    assert [(r["start"], r["end"], r["bold"], r["italic"]) for r in inspected["runs"]] == [
        (0, 3, None, True), (3, 11, None, True), (11, 14, True, None),
    ]
    deleted = apply_docx_edit(data, edit(start=0, end=13, quote="Hello 😀world!", action="replaceText", value=""))
    assert inspect_docx(deleted)["paragraphs"][0]["text"] == ""


@pytest.mark.parametrize("action,value,tag,expected", [("bold", False, "b", "0"), ("italic", True, "i", "1"), ("fontSize", 13.5, "sz", "27")])
def test_formatting_spanning_different_runs_keeps_other_properties(action, value, tag, expected):
    data = package('<w:p><w:r><w:rPr><w:b/><w:color w:val="FF0000"/></w:rPr><w:t>Hel</w:t></w:r><w:r><w:rPr><w:i/></w:rPr><w:t>lo</w:t></w:r></w:p>')
    result = document(apply_docx_edit(data, edit(action=action, value=value)))
    runs = result.xpath("//w:body/w:p/w:r", namespaces=NS)
    assert len(runs) == 2
    assert all(run.find(f"w:rPr/w:{tag}", NS).get(f"{{{W}}}val") == expected for run in runs)
    assert runs[0].find("w:rPr/w:color", NS).get(f"{{{W}}}val") == "FF0000"


def test_paragraph_style_uses_uploaded_definition_and_preserves_section_and_run_formatting():
    data = package('<w:p><w:pPr><w:spacing w:after="70"/><w:sectPr><w:pgMar w:left="999"/></w:sectPr></w:pPr><w:r><w:rPr><w:color w:val="AABBCC"/></w:rPr><w:t>Hello</w:t></w:r></w:p>')
    updated = apply_docx_edit(data, edit(action="style", value="UploadedTitle"))
    result = document(updated)
    assert result.find("w:body/w:p/w:pPr/w:pStyle", NS).get(f"{{{W}}}val") == "UploadedTitle"
    assert result.find("w:body/w:p/w:pPr/w:sectPr/w:pgMar", NS).get(f"{{{W}}}left") == "999"
    assert result.find("w:body/w:p/w:r/w:rPr/w:color", NS).get(f"{{{W}}}val") == "AABBCC"
    assert content(updated)["word/styles.xml"] == content(data)["word/styles.xml"]
    for invalid in ("Heading1", "Emphasis"):
        with pytest.raises(QuickEditError, match="already present"):
            apply_docx_edit(data, edit(action="style", value=invalid))


@pytest.mark.parametrize("inner", [
    '<w:hyperlink><w:r><w:t>Hello</w:t></w:r></w:hyperlink>',
    '<w:ins><w:r><w:t>Hello</w:t></w:r></w:ins>',
    '<w:r><w:rPr><w:rPrChange/></w:rPr><w:t>Hello</w:t></w:r>',
    '<w:r><w:t>Hello</w:t><w:drawing/></w:r>',
    '<w:r><w:t>Hello</w:t><w:tab/></w:r>',
    '<w:bookmarkStart w:id="1"/><w:r><w:t>Hello</w:t></w:r>',
    '<w:pPr><w:pPrChange/></w:pPr><w:r><w:t>Hello</w:t></w:r>',
    '<w:pPr><w:rPr><w:ins/></w:rPr></w:pPr><w:r><w:t>Hello</w:t></w:r>',
    '<w:fldSimple><w:r><w:t>Hello</w:t></w:r></w:fldSimple>',
])
def test_complex_paragraphs_are_readable_but_never_flattened(inner):
    data = package(f"<w:p>{inner}</w:p><w:p><w:r><w:t>Plain</w:t></w:r></w:p>")
    inspected = inspect_docx(data)
    assert not inspected["paragraphs"][0]["editable"]
    assert inspected["paragraphs"][1]["editable"]
    with pytest.raises(QuickEditError, match="unsupported"):
        apply_docx_edit(data, edit())
    assert inspect_docx(apply_docx_edit(data, edit(paragraphId="p:1", quote="Plain")))["paragraphs"][1]["runs"][0]["bold"] is True


def test_field_results_spanning_paragraphs_and_content_controls_cannot_be_edited():
    data = package('''<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r></w:p>
        <w:p><w:r><w:t>Hello</w:t></w:r></w:p>
        <w:p><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>
        <w:sdt><w:sdtContent><w:p><w:r><w:t>Hello</w:t></w:r></w:p></w:sdtContent></w:sdt>
        <w:p><w:r><w:t>Hello</w:t></w:r></w:p>''')
    assert [p["editable"] for p in inspect_docx(data)["paragraphs"]] == [False, False, False, False, True]
    with pytest.raises(QuickEditError):
        apply_docx_edit(data, edit(paragraphId="p:1"))


def test_revision_in_table_row_properties_protects_otherwise_simple_paragraph():
    data = package('<w:tbl><w:tr><w:trPr><w:ins/></w:trPr><w:tc><w:p><w:r><w:t>Hello</w:t></w:r></w:p></w:tc></w:tr></w:tbl>')
    assert not inspect_docx(data)["paragraphs"][0]["editable"]
    with pytest.raises(QuickEditError):
        apply_docx_edit(data, edit())


@pytest.mark.parametrize("settings", ['<w:documentProtection w:enforcement="1"/>', '<w:trackRevisions/>'])
def test_protected_or_tracking_documents_remain_read_only(settings):
    data = package('<w:p><w:r><w:t>Hello</w:t></w:r></w:p>', settings=settings)
    assert not inspect_docx(data)["paragraphs"][0]["editable"]
    with pytest.raises(QuickEditError):
        apply_docx_edit(data, edit())


def test_unicode_offsets_and_repeated_text_use_exact_paragraph_range():
    data = package('<w:p><w:r><w:t>😀Hello Hello</w:t></w:r></w:p>')
    updated = apply_docx_edit(data, edit(start=7, end=12))
    assert [(r["start"], r["end"], r["bold"]) for r in inspect_docx(updated)["paragraphs"][0]["runs"]] == [(0, 7, None), (7, 12, True)]
    with pytest.raises(QuickEditError, match="changed"):
        apply_docx_edit(data, edit(start=2, end=7))  # Wrong UTF-16-derived range.


@pytest.mark.parametrize("changes", [
    {"quote": "World"}, {"start": -1}, {"end": 50}, {"start": True}, {"end": 0},
    {"paragraphId": "p:99"}, {"paragraphId": "p:00"}, {"value": "true"}, {"action": "unknown"},
    {"action": "fontSize", "value": float("nan")}, {"action": "fontSize", "value": 12.1},
    {"action": "fontSize", "value": True}, {"action": "fontSize", "value": 401},
    {"action": "replaceText", "value": "a\nb"}, {"action": "replaceText", "value": "a\tb"},
    {"action": "replaceText", "value": "\ud800"}, {"action": "replaceText", "value": "\u2028"},
    {"action": "replaceText", "value": "a" * 20001},
])
def test_invalid_or_stale_edits_are_rejected(changes):
    data = package('<w:p><w:r><w:t>Hello</w:t></w:r></w:p>')
    with pytest.raises(QuickEditError):
        apply_docx_edit(data, edit(**changes))


@pytest.mark.parametrize("path", ["../escape", "/absolute", "word/../escape", "word\\escape", "word/./escape", "C:drive"])
def test_package_traversal_is_rejected(path):
    with pytest.raises(QuickEditError, match="paths"):
        inspect_docx(package("<w:p/>", extra={path: b"bad"}))


def test_unsafe_xml_duplicates_and_compression_bombs_are_rejected():
    for data in (b"not a zip", package("<w:p/>", styles=b'<!DOCTYPE styles [<!ENTITY payload SYSTEM "file:///etc/passwd">]><styles>&payload;</styles>'), package("<w:p/>", extra={"word/bomb.xml": b"0" * (2 * 1024 * 1024)})):
        with pytest.raises(QuickEditError):
            inspect_docx(data)
    duplicate = BytesIO(package("<w:p/>"))
    with ZipFile(duplicate, "a") as archive, pytest.warns(UserWarning):
        archive.writestr("word/document.xml", b"duplicate")
    with pytest.raises(QuickEditError, match="duplicate"):
        inspect_docx(duplicate.getvalue())
    symlink = BytesIO(package("<w:p/>"))
    with ZipFile(symlink, "a") as archive:
        item = ZipInfo("word/link")
        item.external_attr = 0o120777 << 16
        archive.writestr(item, "../escape")
    with pytest.raises(QuickEditError, match="paths"):
        inspect_docx(symlink.getvalue())


def test_empty_paragraph_before_target_preserves_index_and_original_xml():
    data = package('<w:p/><w:p><w:r><w:t>Hello</w:t></w:r></w:p>')
    updated = apply_docx_edit(data, edit(paragraphId="p:1"))
    assert b"<w:p/>" in content(updated)["word/document.xml"]
    assert inspect_docx(updated)["paragraphs"][1]["runs"][0]["bold"] is True


def test_python_docx_can_reopen_real_uploaded_package_after_edit():
    from docx import Document
    from docx.shared import Pt

    uploaded = Document()
    paragraph = uploaded.add_paragraph()
    run = paragraph.add_run("Hello world")
    run.font.name = "Courier New"
    run.font.size = Pt(12)
    uploaded.sections[0].header.paragraphs[0].text = "Original running header"
    uploaded.add_table(rows=1, cols=1).cell(0, 0).text = "Table unchanged"
    source = BytesIO()
    uploaded.save(source)
    edited = apply_docx_edit(source.getvalue(), edit())
    reopened = Document(BytesIO(edited))
    assert reopened.paragraphs[0].text == "Hello world"
    assert reopened.paragraphs[0].runs[0].bold is True
    assert all(r.font.name == "Courier New" and r.font.size == Pt(12) for r in reopened.paragraphs[0].runs)
    assert reopened.sections[0].header.paragraphs[0].text == "Original running header"
    assert reopened.tables[0].cell(0, 0).text == "Table unchanged"


def test_nonbody_duplicate_text_is_exposed_for_ambiguous_pdf_selection_rejection():
    data = package('<w:p><w:r><w:t>Hello</w:t></w:r></w:p>', extra={
        "word/header2.xml": f'<w:hdr xmlns:w="{W}"><w:p><w:r><w:t>Hello</w:t></w:r></w:p></w:hdr>'.encode(),
        "word/footer1.xml": f'<w:ftr xmlns:w="{W}"><w:p><w:r><w:t>Footer</w:t></w:r></w:p></w:ftr>'.encode(),
        "word/footnotes.xml": f'<w:footnotes xmlns:w="{W}"><w:footnote w:id="1"><w:p><w:r><w:t>Footnote</w:t></w:r></w:p></w:footnote></w:footnotes>'.encode(),
        "word/endnotes.xml": f'<w:endnotes xmlns:w="{W}"><w:endnote w:id="1"><w:p><w:r><w:t>Endnote</w:t></w:r></w:p></w:endnote></w:endnotes>'.encode(),
    })
    result = inspect_docx(data)
    assert result["paragraphs"][0]["id"] == "p:0"
    assert result["paragraphs"][0]["text"] == "Hello" and result["paragraphs"][0]["editable"]
    assert set(result["protectedTexts"]) == {"Hello", "Footer", "Footnote", "Endnote"}
    assert result["hasDynamicFields"] is False
    # There are two possible sources for a "Hello" selection. The UI must not
    # assume that its sole editable body match identifies the selected content.
    assert sum("Hello" in p["text"] for p in result["paragraphs"]) + sum("Hello" in text for text in result["protectedTexts"]) == 2


def test_drawing_chart_math_and_revision_text_are_never_body_selection_targets():
    drawing_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    chart_ns = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    math_ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    data = package(f'''<w:p><w:r><w:t>Chart title</w:t></w:r></w:p>
      <w:p><w:r><w:drawing><a:p xmlns:a="{drawing_ns}"><a:r><a:t>Chart </a:t></a:r><a:r><a:t>title</a:t></a:r></a:p></w:drawing></w:r></w:p>
      <w:p><w:del><w:r><w:delText>Old wording</w:delText></w:r></w:del></w:p>
      <w:p><m:oMath xmlns:m="{math_ns}"><m:r><m:t>x+y</m:t></m:r></m:oMath></w:p>''', extra={
        "word/charts/chart1.xml": f'<c:chartSpace xmlns:c="{chart_ns}"><c:chart><c:val><c:numRef><c:numCache><c:pt idx="0"><c:v>123.5</c:v></c:pt></c:numCache></c:numRef></c:val></c:chart></c:chartSpace>'.encode(),
    })
    result = inspect_docx(data)
    assert [p["editable"] for p in result["paragraphs"]] == [True, False, False, False]
    assert {"Chart title", "Old wording", "x+y", "123.5"}.issubset(result["protectedTexts"])


def test_dynamic_header_fields_are_flagged_and_cached_text_preserved():
    data = package('<w:p><w:r><w:t>Page 2</w:t></w:r></w:p>', extra={
        "word/footer1.xml": f'''<w:ftr xmlns:w="{W}"><w:p><w:r><w:t>Page </w:t></w:r>
          <w:fldSimple w:instr="PAGE"><w:r><w:t>1</w:t></w:r></w:fldSimple></w:p></w:ftr>'''.encode(),
    })
    result = inspect_docx(data)
    assert result["hasDynamicFields"] is True
    assert "Page 1" in result["protectedTexts"]
    # Renderer-updated PAGE results are deliberately not guessed as "Page 2";
    # callers must use selection context to distinguish generated field text.
    assert "Page 2" not in result["protectedTexts"]


def test_unsupported_main_paragraph_text_keeps_visual_tab_separator_for_ambiguity_checks():
    data = package('<w:p><w:r><w:t>Hello</w:t><w:tab/><w:t>world</w:t></w:r></w:p>')
    result = inspect_docx(data)
    assert not result["paragraphs"][0]["editable"]
    assert "Hello world" in result["protectedTexts"]


def test_body_bounds_intersect_custom_section_sizes_and_margins():
    # The fixture's last section is A4 with implicit 1in margins; this earlier
    # section is Letter with a 2in top and 1.5in left margin.
    data = package('''<w:p><w:pPr><w:sectPr><w:pgSz w:w="12240" w:h="15840"/>
      <w:pgMar w:left="2160" w:top="2880" w:right="1440" w:bottom="1440"/></w:sectPr></w:pPr>
      <w:r><w:t>Hello</w:t></w:r></w:p>''')
    bounds = inspect_docx(data)["bodyBounds"]
    assert bounds == pytest.approx({"left": 2160 / 12240, "top": 2880 / 15840,
                                    "right": 1 - 1440 / 11906, "bottom": 1 - 1440 / 15840})


def test_body_bounds_without_section_settings_use_letter_and_one_inch():
    data = package('<w:p><w:r><w:t>Hello</w:t></w:r></w:p>')
    parts = content(data)
    parts["word/document.xml"] = parts["word/document.xml"].replace(b'<w:sectPr><w:pgSz w:w="11906" w:h="16838"/></w:sectPr>', b"")
    output = BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)
    assert inspect_docx(output.getvalue())["bodyBounds"] == pytest.approx({
        "left": 1440 / 12240, "top": 1440 / 15840, "right": 1 - 1440 / 12240, "bottom": 1 - 1440 / 15840,
    })


def test_body_bounds_support_physical_units_mirror_margins_and_gutter():
    data = package('''<w:p><w:pPr><w:sectPr><w:pgSz w:w="8.5in" w:h="11in"/>
      <w:pgMar w:left="1.5in" w:top="1in" w:right="1in" w:bottom="1in" w:gutter="0.25in"/></w:sectPr></w:pPr>
      <w:r><w:t>Hello</w:t></w:r></w:p>''', settings='<w:mirrorMargins/>')
    bounds = inspect_docx(data)["bodyBounds"]
    assert bounds["left"] == pytest.approx(1.75 / 8.5)
    assert bounds["right"] == pytest.approx(1 - 1.75 / 8.5)


@pytest.mark.parametrize("dimension", ['w:w="0" w:h="15840"', 'w:w="invalid" w:h="15840"'])
def test_invalid_page_dimensions_produce_no_safe_body_area(dimension):
    data = package(f'<w:p><w:pPr><w:sectPr><w:pgSz {dimension}/></w:sectPr></w:pPr><w:r><w:t>Hello</w:t></w:r></w:p>')
    assert inspect_docx(data)["bodyBounds"] == {"left": 1.0, "top": 1.0, "right": 0.0, "bottom": 0.0}
