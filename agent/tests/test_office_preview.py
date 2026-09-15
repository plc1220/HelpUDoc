from __future__ import annotations

import asyncio
import base64
import hashlib
import hmac
import io
import json
from pathlib import Path
import shutil
import subprocess
import sys
import time
import zipfile

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
import pytest
from pypdf import PdfReader

from helpudoc_agent.api.routes.office_preview import register_office_preview_routes
from helpudoc_agent.office_preview import (
    OfficePreviewError, OfficePreviewService, convert_to_pdf, decode_content,
    source_revision, validate_office,
)


def docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Quarterly plan", 1)
    document.add_paragraph("Preserve this document's formatting.")
    document.sections[0].header.paragraphs[0].text = "Confidential review"
    document.add_page_break()
    document.add_paragraph("A second page remains a second page.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pptx_bytes() -> bytes:
    presentation = Presentation()
    for title in ("First slide", "Second slide"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = title
    chart = CategoryChartData()
    chart.categories = ["Q1", "Q2"]
    chart.add_series("Revenue", [20, 35])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(7), Inches(3), chart)
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def package_with(content: bytes, name: str, value: bytes) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as original, zipfile.ZipFile(buffer, "w") as output:
        for entry in original.infolist():
            output.writestr(entry, value if entry.filename == name else original.read(entry))
        if name not in original.namelist():
            output.writestr(name, value)
    return buffer.getvalue()


def auth(workspace: str = "workspace-one") -> dict[str, str]:
    encode = lambda content: base64.urlsafe_b64encode(content).decode().rstrip("=")
    header = encode(json.dumps({"alg": "HS256", "typ": "JWT"}).encode())
    payload = encode(json.dumps({"userId": "user-one", "workspaceId": workspace, "exp": time.time() + 60}).encode())
    message = f"{header}.{payload}"
    signature = encode(hmac.new(b"test-secret", message.encode(), hashlib.sha256).digest())
    return {"Authorization": f"Bearer {message}.{signature}"}


def client(converter=lambda *_: b"%PDF-test") -> TestClient:
    app = FastAPI()
    register_office_preview_routes(app, agent_jwt_secret="test-secret", service=OfficePreviewService(converter=converter))
    return TestClient(app)


def request_payload(content=None, filename="review.docx"):
    return {"workspaceId": "workspace-one", "filename": filename,
            "content": base64.b64encode(content or docx_bytes()).decode()}


def test_preview_requires_signed_workspace_context():
    with client() as http:
        payload = request_payload()
        assert http.post("/documents/office-preview", json=payload).status_code == 401
        assert http.post("/documents/office-preview", json=payload, headers=auth("another")).status_code == 403
        result = http.post("/documents/office-preview", json=payload, headers=auth())
        assert result.status_code == 200, result.text
        assert result.json()["document"]["paragraphs"][0]["text"] == "Quarterly plan"
        assert result.json()["revision"] == source_revision(base64.b64decode(payload["content"]))


def test_edit_stale_revision_and_native_formatting():
    content = docx_bytes()
    payload = request_payload(content)
    payload.update({"revision": "0" * 64, "edit": {"paragraphId": "p:1", "start": 0, "end": 8,
                    "quote": "Preserve", "action": "bold", "value": True}})
    with client() as http:
        assert http.post("/documents/office-edit", json=payload, headers=auth()).status_code == 409
        payload["revision"] = source_revision(content)
        result = http.post("/documents/office-edit", json=payload, headers=auth())
        assert result.status_code == 200, result.text
        updated = base64.b64decode(result.json()["content"])
        assert result.json()["revision"] == source_revision(updated)
        with zipfile.ZipFile(io.BytesIO(content)) as before, zipfile.ZipFile(io.BytesIO(updated)) as after:
            for entry in before.namelist():
                if entry != "word/document.xml":
                    assert before.read(entry) == after.read(entry)
        assert Document(io.BytesIO(updated)).paragraphs[1].runs[0].bold is True
        payload["edit"]["quote"] = "Wrong quote"
        assert http.post("/documents/office-edit", json=payload, headers=auth()).status_code == 422


@pytest.mark.parametrize("bad_content", [b"not a zip", b"", b"%PDF-hello"])
def test_invalid_documents_fail_before_converter(bad_content):
    with pytest.raises(OfficePreviewError):
        validate_office(bad_content, "docx")


def test_unsafe_document_packages_are_rejected():
    original = docx_bytes()
    for name, value in [
        ("word/vbaProject.bin", b"macro"),
        ("word/embeddings/oleObject1.bin", b"object"),
        ("../outside", b"unsafe"),
        ("word/document.xml", b'<!DOCTYPE r [<!ENTITY test SYSTEM "file:///etc/passwd">]><r>&test;</r>'),
        ("word/_rels/document.xml.rels", b'<Relationships><Relationship Type="image" TargetMode="External" Target="https://example.com/tracker.png"/></Relationships>'),
        ("word/_rels/document.xml.rels", b'<Relationships><Relationship Type="attachedTemplate" Target="file:///etc/passwd"/></Relationships>'),
        ("word/document.xml", b'<w:document xmlns:w="urn:w"><w:instrText>INCLUDETEXT "/etc/passwd"</w:instrText></w:document>'),
    ]:
        with pytest.raises(OfficePreviewError):
            validate_office(package_with(original, name, value), "docx")


def test_hyperlinks_remain_available_in_document():
    content = package_with(docx_bytes(), "word/_rels/document.xml.rels",
        b'<Relationships><Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" TargetMode="External" Target="https://example.com/"/></Relationships>')
    validate_office(content, "docx")


def test_bad_base64_and_unsupported_format():
    with pytest.raises(OfficePreviewError):
        decode_content("@@@")
    with client() as http:
        assert http.post("/documents/office-preview", json=request_payload(filename="test.docm"), headers=auth()).status_code == 400


def test_same_source_requests_coalesce_and_cache_is_bounded():
    calls = []
    def converter(content, file_format):
        calls.append((content, file_format))
        time.sleep(0.05)
        return b"%PDF-test"
    async def scenario():
        service = OfficePreviewService(converter=converter, cache_bytes=1024 * 1024)
        content = docx_bytes()
        results = await asyncio.gather(*(service.preview("test.docx", content) for _ in range(4)))
        assert len(calls) == 1
        assert results[0] == await service.preview("renamed.docx", content)
        assert len(calls) == 1
        assert not service._pending
        no_cache = OfficePreviewService(converter=converter, cache_bytes=1)
        await no_cache.preview("test.docx", content)
        await no_cache.preview("test.docx", content)
        assert len(calls) == 3
        assert no_cache._cache_size == 0
    asyncio.run(scenario())


def test_missing_converter_is_explicit(monkeypatch):
    monkeypatch.setenv("OFFICE_PREVIEW_SOFFICE", "/does/not/exist")
    with pytest.raises(OfficePreviewError) as error:
        convert_to_pdf(docx_bytes(), "docx")
    assert error.value.status_code == 503


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice not installed")
@pytest.mark.parametrize("file_format,content", [("docx", docx_bytes), ("pptx", pptx_bytes)])
def test_real_office_conversion_retains_pages_and_document_text(file_format, content):
    source = content()
    validate_office(source, file_format)
    pdf = convert_to_pdf(source, file_format)
    reader = PdfReader(io.BytesIO(pdf))
    assert len(reader.pages) == 2
    text = " ".join(page.extract_text() for page in reader.pages)
    if file_format == "docx":
        assert "Quarterly plan" in text
        assert "Confidential review" in text
        assert "second page" in text
    else:
        assert "First slide" in text
        assert "Second slide" in text
        assert "Q1" in text and "Q2" in text


def test_converter_timeout_cleans_up_child(tmp_path, monkeypatch):
    from helpudoc_agent import office_preview
    converter = tmp_path / "slow-converter"
    converter.write_text("#!/bin/sh\nsleep 30\n")
    converter.chmod(0o700)
    monkeypatch.setenv("OFFICE_PREVIEW_SOFFICE", str(converter))
    monkeypatch.setattr(office_preview, "CONVERSION_TIMEOUT", 0.2)
    started = time.monotonic()
    with pytest.raises(OfficePreviewError) as failure:
        convert_to_pdf(docx_bytes(), "docx")
    assert failure.value.status_code == 504
    assert time.monotonic() - started < 5


@pytest.mark.skipif(sys.platform not in ("darwin", "linux"), reason="Supported converter platforms only")
def test_converter_child_cannot_connect_to_network():
    import socket
    from helpudoc_agent import office_preview
    with socket.socket() as listener:
        listener.bind(("127.0.0.1", 0))
        listener.listen(1)
        port = listener.getsockname()[1]
        program = f'''
import socket
try:
    connection = socket.socket()
    connection.settimeout(1)
    connection.connect(("127.0.0.1", {port}))
except PermissionError:
    raise SystemExit(0)
raise SystemExit(1)
'''
        result = subprocess.run([sys.executable, str(Path(office_preview.__file__).resolve()),
                                 "--convert-child", sys.executable, "-c", program],
                                capture_output=True, timeout=5)
        assert result.returncode == 0, result.stderr.decode()


def test_unsafe_nested_and_vector_references_are_rejected():
    for name, value in [
        ("word/_rels/document.xml.rels", b'<Relationships><Relationship Type="image" Target="/etc/passwd"/></Relationships>'),
        ("word/_rels/document.xml.rels", b'<Relationships><Relationship Type="image" Target="../../../../etc/passwd"/></Relationships>'),
        ("word/media/image.svg", b'<svg xmlns="http://www.w3.org/2000/svg"><image href="file:///etc/passwd"/></svg>'),
        ("word/media/image.svg", b'<!DOCTYPE r [<!ENTITY x SYSTEM "file:///etc/passwd">]><r>&x;</r>'),
        ("word/document.xml", b'<w:document xmlns:w="urn:w"><w:altChunk/></w:document>'),
    ]:
        with pytest.raises(OfficePreviewError):
            validate_office(package_with(docx_bytes(), name, value), "docx")


def test_chart_workbooks_render_but_nested_active_content_is_rejected():
    content = pptx_bytes()
    validate_office(content, "pptx")
    with zipfile.ZipFile(io.BytesIO(content)) as package:
        workbook = next(name for name in package.namelist() if name.endswith(".xlsx"))
        macro_workbook = package_with(package.read(workbook), "xl/vbaProject.bin", b"macro")
    with pytest.raises(OfficePreviewError, match="active objects"):
        validate_office(package_with(content, workbook, macro_workbook), "pptx")
    # A package extension cannot make an arbitrary binary an inert workbook.
    with pytest.raises(OfficePreviewError):
        validate_office(package_with(content, workbook, b"not an xlsx"), "pptx")


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice not installed")
def test_real_docx_preview_edit_preview_preserves_uploaded_layout(tmp_path):
    from docx.shared import Pt

    document = Document()
    document.add_heading("Quarterly review", 1)
    document.sections[0].header.paragraphs[0].text = "Confidential / original header"
    paragraph = document.add_paragraph()
    paragraph.add_run("Make ")
    paragraph.add_run("minor ").italic = True
    paragraph.add_run("edits").font.size = Pt(12)
    paragraph.add_run(" with confidence.")
    untouched = "This paragraph and its original layout must remain unchanged."
    document.add_paragraph(untouched)
    table = document.add_table(rows=2, cols=2)
    table.style = "Light Shading Accent 1"
    for cell, text in zip(table.rows[0].cells, ["Owner", "Status"]):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ["Product team", "Ready for review"]):
        cell.text = text
    document.add_page_break()
    document.add_heading("Next steps", 2)
    document.add_paragraph("The second page remains available for annotation.")
    uploaded = tmp_path / "original.docx"
    document.save(uploaded)
    original = uploaded.read_bytes()
    original_revision = source_revision(original)

    async def scenario():
        service = OfficePreviewService()
        before = await service.preview(uploaded.name, original)
        target = next(item for item in before["document"]["paragraphs"] if item["text"] == "Make minor edits with confidence.")
        heading_style = next(item["id"] for item in before["document"]["styles"] if item["name"].lower() == "heading 2")
        current, revision = original, before["revision"]
        selection = {"paragraphId": target["id"], "start": 5, "end": 16, "quote": "minor edits"}
        for action, value in [("bold", True), ("fontSize", 18), ("style", heading_style), ("replaceText", "small updates")]:
            result = await service.edit(uploaded.name, current, revision, {**selection, "action": action, "value": value})
            current = base64.b64decode(result["content"])
            assert result["revision"] != revision
            revision = result["revision"]
        after = await service.preview(uploaded.name, current)
        assert after["revision"] == revision != original_revision
        updated_target = next(item for item in after["document"]["paragraphs"] if item["id"] == target["id"])
        assert updated_target["text"] == "Make small updates with confidence."
        assert updated_target["styleId"] == heading_style
        affected = [run for run in updated_target["runs"] if run["start"] < 18 and run["end"] > 5]
        assert affected and all(run["bold"] is True and run["fontSize"] == 18 for run in affected)
        for item in before["document"]["paragraphs"]:
            if item["id"] != target["id"]:
                assert item == next(other for other in after["document"]["paragraphs"] if other["id"] == item["id"])

        before_pdf = PdfReader(io.BytesIO(base64.b64decode(before["pdf"])))
        after_pdf = PdfReader(io.BytesIO(base64.b64decode(after["pdf"])))
        assert len(before_pdf.pages) == len(after_pdf.pages) == 2
        assert [page.mediabox for page in before_pdf.pages] == [page.mediabox for page in after_pdf.pages]
        before_text = " ".join(page.extract_text() for page in before_pdf.pages)
        after_text = " ".join(page.extract_text() for page in after_pdf.pages)
        assert "minor edits" in before_text and "minor edits" not in after_text
        assert "small updates" in after_text and "small updates" not in before_text
        for expected in ("Confidential / original header", untouched, "Product team", "Ready for review", "Next steps"):
            assert expected in before_text and expected in after_text
        # The PDF itself carries the new text size, beyond the source metadata.
        rendered_runs = []
        after_pdf.pages[0].extract_text(visitor_text=lambda text, _cm, _tm, _font, size: rendered_runs.append((text, size)))
        assert any("small" in text and size == 18 for text, size in rendered_runs)
        with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(io.BytesIO(current)) as edited:
            assert source.namelist() == edited.namelist()
            for name in source.namelist():
                if name != "word/document.xml":
                    assert source.read(name) == edited.read(name)
        # Rendering and quick edits operate on bytes; the uploaded file is untouched.
        assert uploaded.read_bytes() == original
        assert source_revision(uploaded.read_bytes()) == original_revision

    asyncio.run(scenario())
